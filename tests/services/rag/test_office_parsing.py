"""Office document parsing tests for the LlamaIndex RAG pipeline."""

from __future__ import annotations

from pathlib import Path

from deeptutor.services.rag.components.routing import FileTypeRouter
from deeptutor.services.rag.pipelines.llamaindex import LlamaIndexPipeline


def test_file_type_router_routes_docx_and_pptx_to_parser(tmp_path: Path) -> None:
    docx = tmp_path / "面试课程练习.docx"
    pptx = tmp_path / "计划组织专项.pptx"
    legacy_doc = tmp_path / "旧版讲义.doc"
    text = tmp_path / "真题.txt"
    for path in (docx, pptx, legacy_doc, text):
        path.write_bytes(b"placeholder")

    classification = FileTypeRouter.classify_files(
        [str(docx), str(pptx), str(legacy_doc), str(text)]
    )

    assert str(docx) in classification.parser_files
    assert str(pptx) in classification.parser_files
    assert str(legacy_doc) in classification.parser_files
    assert str(text) in classification.text_files


def test_llamaindex_pipeline_extracts_docx_text(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "面试课程练习.docx"
    doc = Document()
    doc.add_paragraph("综合分析题要先表态，再展开原因和对策。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "题型"
    table.cell(0, 1).text = "综合分析"
    doc.save(path)

    pipeline = LlamaIndexPipeline(kb_base_dir=str(tmp_path))
    text = pipeline._extract_docx_text(path)

    assert "综合分析题要先表态" in text
    assert "题型 | 综合分析" in text


def test_llamaindex_pipeline_extracts_pptx_text(tmp_path: Path) -> None:
    from pptx import Presentation

    path = tmp_path / "计划组织专项.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "计划组织题"
    slide.placeholders[1].text = "先明确目标，再拆解流程、资源和风险。"
    presentation.save(path)

    pipeline = LlamaIndexPipeline(kb_base_dir=str(tmp_path))
    text = pipeline._extract_pptx_text(path)

    assert "Slide 1" in text
    assert "计划组织题" in text
    assert "先明确目标" in text


def test_llamaindex_pipeline_extracts_legacy_doc_text_with_converter(
    monkeypatch,
    tmp_path: Path,
) -> None:
    path = tmp_path / "旧版讲义.doc"
    path.write_bytes(b"legacy doc placeholder")
    converted_text = "论证逻辑要围绕问题、原因、意义和对策展开。"

    def fake_run(cmd, check, capture_output, text, timeout):
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        (outdir / "旧版讲义.txt").write_text(converted_text, encoding="utf-8")
        return object()

    monkeypatch.setattr(
        LlamaIndexPipeline,
        "_find_office_converter",
        staticmethod(lambda: "libreoffice"),
    )
    monkeypatch.setattr(
        "deeptutor.services.rag.pipelines.llamaindex.subprocess.run",
        fake_run,
    )

    pipeline = LlamaIndexPipeline(kb_base_dir=str(tmp_path))
    text = pipeline._extract_legacy_doc_text(path)

    assert converted_text in text
